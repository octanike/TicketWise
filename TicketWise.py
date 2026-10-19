#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GINEFOBIA - Registro de entradas
---------------------------------
Aplicacion de escritorio (Tkinter) para registrar entradas vendidas
y exportarlas a Word / Excel. Los datos se guardan en una base de
datos SQLite local (ginefobia.db), en la misma carpeta que este archivo.

Requisitos:
    pip install python-docx openpyxl

Ejecutar:
    python ginefobia.py
"""

import os
import sys
import sqlite3
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import date

# ---------------------------------------------------------------------------
# Chequeo de dependencias (con mensaje amigable en vez de un traceback crudo)
# ---------------------------------------------------------------------------
try:
    import docx  # noqa: F401  (python-docx)
    import openpyxl  # noqa: F401
except ImportError as _imp_err:
    print("Faltan paquetes necesarios para exportar a Word/Excel.")
    print("Instalalos con:\n")
    print("    pip install python-docx openpyxl\n")
    print(f"Detalle: {_imp_err}")
    sys.exit(1)


# ---------------------------------------------------------------------------
# Configuracion / paleta de colores
# ---------------------------------------------------------------------------

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "ginefobia.db")

PAPER = "#ECEFE6"
PAPER_CARD = "#F7F8F2"
PAPER_CARD_ALT = "#EFF1E9"
INK = "#1E2420"
INK_MUTED = "#5B6459"
RED = "#A8362A"
RED_DEEP = "#7C2620"
BLUE = "#35506B"
LINE = "#C6CBBC"

FONT_TITLE = ("Consolas", 34, "bold")
FONT_MONO_SM = ("Consolas", 8)
FONT_MONO_MD = ("Consolas", 9)
FONT_MONO_BOLD = ("Consolas", 10, "bold")
FONT_HEADING = ("Segoe UI", 16, "bold")
FONT_LABEL = ("Segoe UI", 9)
FONT_BODY = ("Segoe UI", 10)
FONT_CARD_TITLE = ("Segoe UI", 13, "bold")
FONT_CARD_MARK = ("Consolas", 26, "bold")

METODOS = ["Efectivo", "Transferencia", "Debito", "Credito", "Mercado Pago"]


# ---------------------------------------------------------------------------
# Base de datos (SQLite)
# ---------------------------------------------------------------------------

def init_db():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS entradas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL,
            apellido TEXT NOT NULL,
            cantidad INTEGER NOT NULL,
            precio REAL NOT NULL,
            metodo TEXT NOT NULL,
            fecha TEXT NOT NULL
        )
        """
    )
    conn.commit()
    conn.close()


def agregar_entrada(nombre, apellido, cantidad, precio, metodo, fecha):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO entradas (nombre, apellido, cantidad, precio, metodo, fecha) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (nombre, apellido, cantidad, precio, metodo, fecha),
    )
    conn.commit()
    conn.close()


def obtener_entradas():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()
    cur.execute("SELECT * FROM entradas ORDER BY id ASC")
    filas = [dict(r) for r in cur.fetchall()]
    conn.close()
    return filas


def contar_entradas():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM entradas")
    n = cur.fetchone()[0]
    conn.close()
    return n


def borrar_todas_entradas():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("DELETE FROM entradas")
    cur.execute("DELETE FROM sqlite_sequence WHERE name='entradas'")
    conn.commit()
    conn.close()


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def formatear_moneda(valor):
    try:
        valor = float(valor)
    except (TypeError, ValueError):
        valor = 0.0
    texto = f"{valor:,.2f}"          # "1,234.50" (formato US)
    entero, decimales = texto.split(".")
    entero = entero.replace(",", ".")  # "1.234"
    return f"$ {entero},{decimales}"   # "$ 1.234,50"


# ---------------------------------------------------------------------------
# Exportacion a Word / Excel
# ---------------------------------------------------------------------------

def generar_docx(entradas, ruta):
    from docx import Document
    from docx.shared import RGBColor

    doc = Document()

    titulo = doc.add_heading("GINEFOBIA", level=0)
    for run in titulo.runs:
        run.font.color.rgb = RGBColor(0xA8, 0x36, 0x2A)

    doc.add_heading("Listado de entradas", level=2)

    encabezados = ["Nombre", "Apellido", "Cantidad", "Precio", "Metodo", "Fecha", "Total"]
    table = doc.add_table(rows=1, cols=len(encabezados))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(encabezados):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True

    total_general = 0.0
    for e in entradas:
        total = e["cantidad"] * e["precio"]
        total_general += total
        row_cells = table.add_row().cells
        row_cells[0].text = e["nombre"]
        row_cells[1].text = e["apellido"]
        row_cells[2].text = str(e["cantidad"])
        row_cells[3].text = formatear_moneda(e["precio"])
        row_cells[4].text = e["metodo"]
        row_cells[5].text = e["fecha"]
        row_cells[6].text = formatear_moneda(total)

    doc.add_paragraph()
    doc.add_paragraph().add_run(f"Total de entradas: {len(entradas)}").bold = True
    doc.add_paragraph().add_run(f"Total recaudado: {formatear_moneda(total_general)}").bold = True

    doc.save(ruta)


def generar_xlsx(entradas, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Entradas"

    encabezados = ["Nombre", "Apellido", "Cantidad", "Precio", "Metodo", "Fecha", "Total"]
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True)

    for e in entradas:
        total = e["cantidad"] * e["precio"]
        ws.append([e["nombre"], e["apellido"], e["cantidad"], e["precio"], e["metodo"], e["fecha"], total])

    anchos = [16, 16, 10, 12, 16, 12, 12]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[chr(64 + i)].width = ancho

    wb.save(ruta)


# ---------------------------------------------------------------------------
# Interfaz grafica
# ---------------------------------------------------------------------------

class GinefobiaApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("GINEFOBIA - Registro de entradas")
        self.geometry("900x640")
        self.minsize(780, 540)
        self.configure(background=PAPER)

        self._configurar_estilos()

        container = tk.Frame(self, background=PAPER)
        container.pack(fill="both", expand=True)
        container.grid_rowconfigure(0, weight=1)
        container.grid_columnconfigure(0, weight=1)

        self.frames = {}
        for Clase, nombre in (
            (LobbyFrame, "lobby"),
            (NuevaEntradaFrame, "nueva"),
            (EntradasFrame, "entradas"),
        ):
            frame = Clase(container, self)
            self.frames[nombre] = frame
            frame.grid(row=0, column=0, sticky="nsew")

        self.mostrar("lobby")

    def _configurar_estilos(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        style.configure("TFrame", background=PAPER)
        style.configure("Card.TFrame", background=PAPER_CARD)

        style.configure("TLabel", background=PAPER, foreground=INK, font=FONT_BODY)
        style.configure("Card.TLabel", background=PAPER_CARD, foreground=INK, font=FONT_BODY)
        style.configure("Muted.TLabel", background=PAPER, foreground=INK_MUTED, font=FONT_LABEL)
        style.configure("CardMuted.TLabel", background=PAPER_CARD, foreground=INK_MUTED, font=FONT_MONO_SM)
        style.configure("Heading.TLabel", background=PAPER, foreground=INK, font=FONT_HEADING)

        style.configure("TEntry", fieldbackground="white", foreground=INK, padding=5)
        style.configure("TCombobox", fieldbackground="white", foreground=INK, padding=4)
        style.configure("TSpinbox", fieldbackground="white", foreground=INK, padding=4, arrowsize=14)

        style.configure(
            "Primary.TButton", background=RED, foreground="white",
            font=FONT_MONO_BOLD, padding=(18, 10), borderwidth=0,
        )
        style.map(
            "Primary.TButton",
            background=[("disabled", LINE), ("active", RED_DEEP)],
            foreground=[("disabled", INK_MUTED)],
        )

        style.configure(
            "Outline.TButton", background=PAPER_CARD, foreground=INK,
            font=FONT_BODY, padding=(14, 8), borderwidth=1, relief="solid",
        )
        style.map(
            "Outline.TButton",
            background=[("disabled", PAPER_CARD), ("active", BLUE)],
            foreground=[("disabled", LINE), ("active", "white")],
        )

        style.configure(
            "Danger.TButton", background=PAPER_CARD, foreground=RED,
            font=FONT_BODY, padding=(14, 8), borderwidth=1, relief="solid",
        )
        style.map(
            "Danger.TButton",
            background=[("disabled", PAPER_CARD), ("active", RED)],
            foreground=[("disabled", LINE), ("active", "white")],
        )

        style.configure(
            "Link.TButton", background=PAPER, foreground=INK_MUTED,
            font=FONT_MONO_MD, borderwidth=0, padding=2,
        )
        style.map("Link.TButton", foreground=[("active", INK)])

    def mostrar(self, nombre):
        frame = self.frames[nombre]
        if nombre == "entradas":
            frame.refrescar()
        elif nombre == "lobby":
            frame.actualizar_contador()
        frame.tkraise()


class LobbyFrame(ttk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, style="TFrame")
        self.controller = controller

        wrap = ttk.Frame(self, style="TFrame")
        wrap.place(relx=0.5, rely=0.5, anchor="center")

        tk.Label(
            wrap, text="GINEFOBIA", font=FONT_TITLE, background=PAPER, foreground=RED,
        ).pack(pady=(0, 4))

        ttk.Label(
            wrap, text="CONTROL DE ENTRADAS", style="Muted.TLabel", font=FONT_MONO_MD,
        ).pack(pady=(0, 18))

        tk.Frame(wrap, background=INK, height=2, width=520).pack(pady=(0, 30))

        cards = ttk.Frame(wrap, style="TFrame")
        cards.pack()

        self.card_nueva, _ = self._crear_card(
            cards, "+", "Nueva Entrada", "Registrar una venta", RED,
            lambda: controller.mostrar("nueva"),
        )
        self.card_nueva.grid(row=0, column=0, padx=12)

        self.card_entradas, self.lbl_entradas_sub = self._crear_card(
            cards, "=", "Entradas", "Cargando...", BLUE,
            lambda: controller.mostrar("entradas"),
        )
        self.card_entradas.grid(row=0, column=1, padx=12)

    def _crear_card(self, parent, marca, titulo, subtitulo_inicial, color, comando):
        card = tk.Frame(
            parent, background=PAPER_CARD, highlightbackground=LINE,
            highlightthickness=1, cursor="hand2",
        )
        borde = tk.Frame(card, background=color, width=4)
        borde.pack(side="left", fill="y")

        contenido = tk.Frame(card, background=PAPER_CARD, cursor="hand2")
        contenido.pack(side="left", fill="both", expand=True, padx=24, pady=20)

        lbl_marca = tk.Label(
            contenido, text=marca, font=FONT_CARD_MARK, background=PAPER_CARD, foreground=color,
        )
        lbl_marca.pack(anchor="w")
        lbl_titulo = tk.Label(
            contenido, text=titulo, font=FONT_CARD_TITLE, background=PAPER_CARD, foreground=INK,
        )
        lbl_titulo.pack(anchor="w", pady=(10, 0))
        lbl_sub = tk.Label(
            contenido, text=subtitulo_inicial, font=FONT_LABEL, background=PAPER_CARD, foreground=INK_MUTED,
        )
        lbl_sub.pack(anchor="w")

        widgets = [card, borde, contenido, lbl_marca, lbl_titulo, lbl_sub]
        for w in widgets:
            w.bind("<Button-1>", lambda e: comando())

        def al_entrar(_e):
            card.configure(highlightbackground=color, highlightthickness=2)

        def al_salir(_e):
            card.configure(highlightbackground=LINE, highlightthickness=1)

        for w in widgets:
            w.bind("<Enter>", al_entrar)
            w.bind("<Leave>", al_salir)

        return card, lbl_sub

    def actualizar_contador(self):
        n = contar_entradas()
        palabra = "registrada" if n == 1 else "registradas"
        self.lbl_entradas_sub.configure(text=f"{n} {palabra}")


class NuevaEntradaFrame(ttk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, style="TFrame")
        self.controller = controller

        top = ttk.Frame(self, style="TFrame")
        top.pack(fill="x", padx=40, pady=(28, 6))
        ttk.Button(
            top, text="<  Volver", style="Link.TButton",
            command=lambda: controller.mostrar("lobby"),
        ).pack(anchor="w")

        wrap = ttk.Frame(self, style="TFrame")
        wrap.pack(fill="both", expand=True, padx=40, pady=(4, 20))

        ttk.Label(wrap, text="Nueva Entrada", style="Heading.TLabel").pack(anchor="w")
        ttk.Label(
            wrap, text="Completa los datos de la venta.", style="Muted.TLabel",
        ).pack(anchor="w", pady=(2, 18))

        self.lbl_error_guardado = tk.Label(
            wrap, text="", background=PAPER, foreground=RED, font=FONT_LABEL,
        )
        self.lbl_error_guardado.pack(anchor="w")

        card = tk.Frame(wrap, background=PAPER_CARD, highlightbackground=LINE, highlightthickness=1)
        card.pack(fill="x")

        form = ttk.Frame(card, style="Card.TFrame")
        form.pack(fill="x", padx=30, pady=(24, 6))
        form.columnconfigure(0, weight=1, uniform="col")
        form.columnconfigure(1, weight=1, uniform="col")

        self.var_nombre = tk.StringVar()
        self.var_apellido = tk.StringVar()
        self._campo_texto(form, "NOMBRE", self.var_nombre, 0, 1, 0)
        self._campo_texto(form, "APELLIDO", self.var_apellido, 0, 1, 1)

        self.var_cantidad = tk.StringVar(value="1")
        self.var_precio = tk.StringVar()
        self._campo_spinbox(form, "CANTIDAD", self.var_cantidad, 2, 3, 0, 1, 999)
        self._campo_texto(form, "PRECIO", self.var_precio, 2, 3, 1, placeholder="0.00")

        ttk.Label(form, text="METODO DE PAGO", style="CardMuted.TLabel").grid(
            row=4, column=0, sticky="w", pady=(0, 4)
        )
        self.var_metodo = tk.StringVar(value=METODOS[0])
        ttk.Combobox(
            form, textvariable=self.var_metodo, values=METODOS, state="readonly",
        ).grid(row=5, column=0, sticky="ew", padx=(0, 15), pady=(0, 18))

        ttk.Label(form, text="FECHA (DD/MM/AA)", style="CardMuted.TLabel").grid(
            row=4, column=1, sticky="w", pady=(0, 4)
        )
        fecha_frame = ttk.Frame(form, style="Card.TFrame")
        fecha_frame.grid(row=5, column=1, sticky="w", pady=(0, 18))

        hoy = date.today()
        self.var_dia = tk.StringVar(value=f"{hoy.day:02d}")
        self.var_mes = tk.StringVar(value=f"{hoy.month:02d}")
        self.var_anio = tk.StringVar(value=f"{hoy.year % 100:02d}")

        ttk.Spinbox(
            fecha_frame, from_=1, to=31, width=3, textvariable=self.var_dia,
            format="%02.0f", wrap=True, justify="center",
        ).pack(side="left")
        ttk.Label(fecha_frame, text="/", style="CardMuted.TLabel", font=FONT_BODY).pack(side="left", padx=5)
        ttk.Spinbox(
            fecha_frame, from_=1, to=12, width=3, textvariable=self.var_mes,
            format="%02.0f", wrap=True, justify="center",
        ).pack(side="left")
        ttk.Label(fecha_frame, text="/", style="CardMuted.TLabel", font=FONT_BODY).pack(side="left", padx=5)
        ttk.Spinbox(
            fecha_frame, from_=20, to=50, width=3, textvariable=self.var_anio,
            format="%02.0f", wrap=True, justify="center",
        ).pack(side="left")

        acciones = ttk.Frame(card, style="Card.TFrame")
        acciones.pack(fill="x", padx=30, pady=(4, 24))

        ttk.Button(
            acciones, text="GUARDAR", style="Primary.TButton", command=self.guardar,
        ).pack(side="left")

        self.lbl_estado = tk.Label(
            acciones, text="", background=PAPER_CARD, foreground=RED, font=FONT_MONO_BOLD,
        )
        self.lbl_estado.pack(side="left", padx=16)

    def _campo_texto(self, parent, etiqueta, var, fila_label, fila_campo, col, placeholder=None):
        padx = (0, 15) if col == 0 else (0, 0)
        ttk.Label(parent, text=etiqueta, style="CardMuted.TLabel").grid(
            row=fila_label, column=col, sticky="w", pady=(0, 4)
        )
        entry = ttk.Entry(parent, textvariable=var)
        entry.grid(row=fila_campo, column=col, sticky="ew", padx=padx, pady=(0, 18))

    def _campo_spinbox(self, parent, etiqueta, var, fila_label, fila_campo, col, desde, hasta):
        padx = (0, 15) if col == 0 else (0, 0)
        ttk.Label(parent, text=etiqueta, style="CardMuted.TLabel").grid(
            row=fila_label, column=col, sticky="w", pady=(0, 4)
        )
        sp = ttk.Spinbox(parent, from_=desde, to=hasta, textvariable=var, justify="center")
        sp.grid(row=fila_campo, column=col, sticky="ew", padx=padx, pady=(0, 18))

    def guardar(self):
        self.lbl_error_guardado.configure(text="")
        nombre = self.var_nombre.get().strip()
        apellido = self.var_apellido.get().strip()

        if not nombre or not apellido:
            messagebox.showwarning("Datos incompletos", "Completa nombre y apellido.")
            return

        try:
            cantidad = int(self.var_cantidad.get())
            if cantidad <= 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Datos incompletos", "La cantidad tiene que ser un numero entero mayor a 0.")
            return

        try:
            precio = float(self.var_precio.get().strip().replace(",", "."))
            if precio < 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Datos incompletos", "El precio tiene que ser un numero valido.")
            return

        try:
            dia = int(self.var_dia.get())
            mes = int(self.var_mes.get())
            anio = int(self.var_anio.get())
            fecha = f"{dia:02d}/{mes:02d}/{anio:02d}"
        except ValueError:
            messagebox.showwarning("Datos incompletos", "Revisa la fecha.")
            return

        agregar_entrada(nombre, apellido, cantidad, precio, self.var_metodo.get(), fecha)

        self.var_nombre.set("")
        self.var_apellido.set("")
        self.var_cantidad.set("1")
        self.var_precio.set("")
        self.var_metodo.set(METODOS[0])

        self.lbl_estado.configure(text="ENTRADA REGISTRADA \u2713")
        self.after(3000, lambda: self.lbl_estado.configure(text=""))


class EntradasFrame(ttk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, style="TFrame")
        self.controller = controller

        top = ttk.Frame(self, style="TFrame")
        top.pack(fill="x", padx=40, pady=(28, 6))
        ttk.Button(
            top, text="<  Volver", style="Link.TButton",
            command=lambda: controller.mostrar("lobby"),
        ).pack(anchor="w")

        header = ttk.Frame(self, style="TFrame")
        header.pack(fill="x", padx=40, pady=(4, 12))
        ttk.Label(header, text="Entradas", style="Heading.TLabel").pack(side="left")
        self.lbl_resumen = tk.Label(
            header, text="", background=PAPER, foreground=INK_MUTED, font=FONT_MONO_MD,
        )
        self.lbl_resumen.pack(side="right")

        cont_tabla = tk.Frame(self, background=PAPER_CARD, highlightbackground=LINE, highlightthickness=1)
        cont_tabla.pack(fill="both", expand=True, padx=40, pady=(0, 12))

        columnas = ("nombre", "apellido", "cantidad", "precio", "metodo", "fecha", "total")
        self.tree = ttk.Treeview(cont_tabla, columns=columnas, show="headings", height=12)
        encabezados = {
            "nombre": "Nombre", "apellido": "Apellido", "cantidad": "Cant.",
            "precio": "Precio", "metodo": "Metodo", "fecha": "Fecha", "total": "Total",
        }
        anchos = {"nombre": 130, "apellido": 130, "cantidad": 60, "precio": 100, "metodo": 130, "fecha": 80, "total": 110}
        for c in columnas:
            self.tree.heading(c, text=encabezados[c])
            ancla = "e" if c in ("cantidad", "precio", "total") else "w"
            self.tree.column(c, width=anchos[c], anchor=ancla)

        self.scrollbar = ttk.Scrollbar(cont_tabla, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=self.scrollbar.set)
        # El tree y la scrollbar se empaquetan/desempaquetan desde refrescar(),
        # segun haya o no entradas, para poder mostrar el mensaje de "vacio".

        self.tree.tag_configure("par", background=PAPER_CARD)
        self.tree.tag_configure("impar", background=PAPER_CARD_ALT)

        self.lbl_vacio = tk.Label(
            cont_tabla, text="Todavia no hay entradas registradas.",
            background=PAPER_CARD, foreground=INK_MUTED, font=FONT_BODY,
        )

        acciones = ttk.Frame(self, style="TFrame")
        acciones.pack(fill="x", padx=40, pady=(0, 28))

        self.btn_word = ttk.Button(
            acciones, text="Imprimir en Word (.docx)", style="Outline.TButton", command=self.exportar_word,
        )
        self.btn_word.pack(side="left")

        self.btn_excel = ttk.Button(
            acciones, text="Imprimir en Excel (.xlsx)", style="Outline.TButton", command=self.exportar_excel,
        )
        self.btn_excel.pack(side="left", padx=10)

        self.btn_borrar = ttk.Button(
            acciones, text="Borrar todo", style="Danger.TButton", command=self.borrar_todo,
        )
        self.btn_borrar.pack(side="right")

    def refrescar(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        entradas = obtener_entradas()
        total_general = 0.0
        for i, e in enumerate(entradas):
            total = e["cantidad"] * e["precio"]
            total_general += total
            tag = "par" if i % 2 == 0 else "impar"
            self.tree.insert(
                "", "end",
                values=(
                    e["nombre"], e["apellido"], e["cantidad"],
                    formatear_moneda(e["precio"]), e["metodo"], e["fecha"],
                    formatear_moneda(total),
                ),
                tags=(tag,),
            )

        if entradas:
            self.lbl_vacio.pack_forget()
            self.tree.pack(side="left", fill="both", expand=True, padx=(1, 0), pady=1)
            self.scrollbar.pack(side="right", fill="y")
            self.lbl_resumen.configure(
                text=f"{len(entradas)} entradas   \u00b7   Total {formatear_moneda(total_general)}"
            )
        else:
            self.tree.pack_forget()
            self.scrollbar.pack_forget()
            self.lbl_vacio.pack(expand=True)
            self.lbl_resumen.configure(text="")

        estado = "!disabled" if entradas else "disabled"
        for boton in (self.btn_word, self.btn_excel, self.btn_borrar):
            boton.state([estado])

    def exportar_word(self):
        entradas = obtener_entradas()
        if not entradas:
            return
        nombre_sugerido = f"entradas_ginefobia_{date.today().strftime('%Y%m%d')}.docx"
        ruta = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Documento de Word", "*.docx")],
            initialfile=nombre_sugerido,
            title="Guardar listado como",
        )
        if not ruta:
            return
        try:
            generar_docx(entradas, ruta)
            messagebox.showinfo("Listo", f"Se exporto el archivo:\n{ruta}")
        except Exception as ex:
            messagebox.showerror("Error", f"No se pudo exportar a Word.\n\n{ex}")

    def exportar_excel(self):
        entradas = obtener_entradas()
        if not entradas:
            return
        nombre_sugerido = f"entradas_ginefobia_{date.today().strftime('%Y%m%d')}.xlsx"
        ruta = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Libro de Excel", "*.xlsx")],
            initialfile=nombre_sugerido,
            title="Guardar listado como",
        )
        if not ruta:
            return
        try:
            generar_xlsx(entradas, ruta)
            messagebox.showinfo("Listo", f"Se exporto el archivo:\n{ruta}")
        except Exception as ex:
            messagebox.showerror("Error", f"No se pudo exportar a Excel.\n\n{ex}")

    def borrar_todo(self):
        entradas = obtener_entradas()
        if not entradas:
            return
        confirmar = messagebox.askyesno(
            "Borrar todas las entradas",
            f"Se van a eliminar las {len(entradas)} entradas registradas de forma "
            "permanente.\n\nEsta accion no se puede deshacer. \u00bfQueres continuar?",
            icon="warning",
        )
        if confirmar:
            borrar_todas_entradas()
            self.refrescar()


def main():
    init_db()
    app = GinefobiaApp()
    app.mainloop()


if __name__ == "__main__":
    main()
